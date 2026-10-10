from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION_START
from scipy.stats import norm
import math
from pathlib import Path

out = Path('/mnt/data/Citi_Coffee_Risk_Management_Analysis.docx')
S=1.2; r=.02; u=.01; T=.5; K=1.25; sigma=.25; qty=100000
F=S*math.exp((r+u)*T)
d1=(math.log(F/K)+(sigma*sigma/2)*T)/(sigma*math.sqrt(T));d2=d1-sigma*math.sqrt(T)
call=math.exp(-r*T)*(F*norm.cdf(d1)-K*norm.cdf(d2))
var=qty*S*1.645*sigma*math.sqrt(1/12)
navy=RGBColor(23,43,66); blue=RGBColor(17,93,140); teal=RGBColor(11,112,128); gray=RGBColor(96,109,124)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill);tcPr.append(sh)
def set_border(cell,bottom='D9E0E7'):
    tcPr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders'); e=OxmlElement('w:bottom');e.set(qn('w:val'),'single'); e.set(qn('w:color'),bottom);e.set(qn('w:sz'),'4'); borders.append(e); tcPr.append(borders)
def add_bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.add_run(text); return p
def add_heading(doc,title,level=2):
    doc.add_heading(title,level)
def add_table(doc,heads,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid'
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i];c.text=h;shade(c,'17304B')
        for p in c.paragraphs:
            for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(9)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            cells[i].text=str(x);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j%2==1:shade(cells[i],'F0F5F9')
            for p in cells[i].paragraphs:
                for run in p.runs:run.font.size=Pt(8.5)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):row.cells[i].width=Inches(w)
    return t

d=Document(); s=d.sections[0];s.top_margin=Inches(.65);s.bottom_margin=Inches(.59);s.left_margin=Inches(.72);s.right_margin=Inches(.72)
styles=d.styles; styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(8.8);styles['Normal'].font.color.rgb=navy
styles['Normal'].paragraph_format.space_after=Pt(3.5)
for name,size,color in [('Title',21,navy),('Heading 1',14,navy),('Heading 2',10.5,blue)]:
    st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=color
    st.paragraph_format.space_before=Pt(8 if name!='Title' else 0);st.paragraph_format.space_after=Pt(3.5)

header=s.header.paragraphs[0];header.text='CITI MQA  |  LEARNING SIMULATION';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r1 in header.runs:r1.font.size=Pt(8);r1.font.bold=True;r1.font.color.rgb=gray
footer=s.footer.paragraphs[0];footer.text='Prepared for Citi / Forage virtual experience  •  Illustrative educational analysis';footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
for ru in footer.runs: ru.font.size=Pt(7.4);ru.font.color.rgb=gray

p=d.add_paragraph(style='Title');p.add_run('Coffee Futures: Risk Management Analysis')
p=d.add_paragraph();p.add_run('Market, Credit and Operational Risk  |  Three Client Scenarios').bold=True
p=d.add_paragraph('Decision-oriented proposal for the Markets Quantitative Analysis training exercise.');p.style='Subtitle'

add_heading(d,'Executive summary',1)
d.add_paragraph('Scenario 1: For a coffee purchaser exposed to rising prices, buy coffee futures to approximately lock the acquisition cost or buy call options to cap price increases while retaining benefits from lower prices. Scenario 2: Onboard a new supplier through credit due diligence, enforceable delivery protections and collateral or bank guarantees. Scenario 3: Reduce manual execution risk through controlled automation, pre-trade validations, reconciliation and incident governance.')

add_heading(d,'Scenario 1 | Hedging market risk',1)
p=d.add_paragraph();p.add_run('Exposure and decision. ').bold=True;p.add_run('Assume a business will purchase 100,000 lb of coffee in six months. An adverse-weather supply shock could increase procurement cost. Compare a long futures hedge with a long call on coffee futures.')
add_table(d,['Model input','Value','Model input','Value'],[
('Coffee spot price','$1.20 / lb','Annual risk-free rate','2.00%'),('Storage carry rate','1.00%','Time to maturity','0.50 years'),('Option strike','$1.25 / lb','Annual volatility','25.00%')],[1.55,1.65,1.65,1.6])
p=d.add_paragraph();p.add_run('Quantitative pricing (illustrative). ').bold=True;p.add_run(f'Cost of carry F(0,T) = S(0) × exp[(r + u)T] = ${F:.5f}/lb. An at-the-maturity European call on the coffee futures contract, valued using Black-76 at volatility 25%, is ${call:.5f}/lb (approximately ${call*qty:,.0f} premium for 100,000 lb).')
p=d.add_paragraph();p.add_run('Alternative hedges. ').bold=True;p.add_run('A long futures position offsets higher physical purchase costs but creates variation-margin demands and foregoes lower-price savings. A purchased call retains lower-price benefits but has an upfront premium and a strike-dependent cap on costs. Use the appropriate exchange contract quantity and avoid over-hedging; the scenario below assumes perfect matching at maturity.')
rows=[]
for price in [.96,1.20,1.44]:
    physical=price*qty; netf=F*qty; netc=(price-max(price-K,0)+call)*qty
    rows.append((f'${price:.2f}',f'${physical:,.0f}',f'${netf:,.0f}',f'${netc:,.0f}'))
add_table(d,['Maturity spot ($/lb)','No hedge ($)','Long futures ($)','Long call ($)'],rows,[1.55,1.6,1.6,1.6])
p=d.add_paragraph('Table: Total illustrative coffee procurement cost for 100,000 lb; long-call column includes the Black-76 premium. Assumes matched futures/physical prices at expiry and excludes basis, execution fees, collateral funding and counterparty effects.');p.style='Caption'
add_bullet(d,f'Risk measurement: Unhedged one-month 95% parametric VaR ≈ ${var:,.0f} [100,000 × $1.20 × 1.645 × 25% × sqrt(1/12)], under a normal-return approximation; exceedances can be larger.')
add_bullet(d,'Stress tests: reprice for +10%, +20%, +30% coffee moves, a volatility spike, a 5% unfavorable basis change and stressed margin demands; record the effect on liquidity and coverage.')
add_bullet(d,'Monitoring: daily hedge ratio, futures basis, collateral availability, volatility, forecast production disruptions and hedge effectiveness; rebalance if exposure or timing changes.')

add_heading(d,'Scenario 2 | Credit risk with a new supplier',1)
p=d.add_paragraph();p.add_run('Exposure and decision. ').bold=True;p.add_run('A supplier in a volatile economy may fail to deliver, require costly replacement purchases or expose advances to default. Credit protection must cover performance/non-delivery risks, not merely financial payment defaults.')
add_bullet(d,'Onboarding: verify ownership and financial statements, payment/delivery history, sanctions compliance, country/transfer risk, concentration and business continuity. Use rating analysis, logistic regression and relevant ratios (Altman Z-score only where appropriate to the company type).')
add_bullet(d,'Quantification: use PD (probability of default), LGD (loss given default) and EAD (exposure at default): Expected Loss = PD × LGD × EAD. For an illustrative $120,000 exposure, PD 5% and LGD 60%, expected loss is $3,600. These are assumptions, not assessed supplier statistics.')
add_bullet(d,'Mitigation: minimize prepayments; require a strong-bank letter of credit, performance bond or secured escrow as appropriate, with clear delivery milestones, inspection rights, collateral eligibility and haircuts; diversify sourcing and set per-counterparty/country limits.')
add_bullet(d,'Monitoring: review financial/credit changes, missed milestones, collateral sufficiency, exposure limits, political developments and early-warning triggers. Escalate deteriorating scores; pause or reduce new orders when thresholds are breached.')
p=d.add_paragraph();p.add_run('CDS caveat. ').bold=True;p.add_run('A credit default swap may be unavailable for a private physical supplier and is not a direct substitute for a performance guarantee on physical coffee delivery.')

add_heading(d,'Scenario 3 | Operational risk in trade execution',1)
p=d.add_paragraph();p.add_run('Exposure and decision. ').bold=True;p.add_run('Manual order entry creates wrong-side trades, oversized positions, duplicate tickets, delayed confirmations and inaccurate risk reporting. The priority is preventive controls plus prompt detection/recovery.')
add_bullet(d,'Automate a single order workflow with maker-checker approval for exceptional orders, unique order IDs, time synchronization, immutable logs and restricted permissions.')
add_bullet(d,'Introduce pre-trade hard controls for allowed instruments, trader permissions, maximum order size/notional, price collars, credit/margin headroom, daily loss limits and duplicate-order prevention; provide a kill switch.')
add_bullet(d,'Reconcile order → execution → clearing/broker statements → positions/P&L intraday and at end of day; alert on unmatched trades, rejects, late fills and stale market data.')
add_bullet(d,'Implement system testing, user-acceptance testing, fallback/rollback, access reviews, cybersecurity, disaster recovery and staff training. Run incident postmortems with assigned remediation owners.')
add_bullet(d,'Track error rate per 1,000 tickets, number of reconciliation breaks, false rejects, time to detection/recovery and percent of trades passing automatic controls. Monitor trends and impose escalation thresholds.')

add_heading(d,'Risk ownership and control cadence',1)
add_table(d,['Risk','Primary controls','Monitoring / escalation'],[
('Market','Matched futures or call hedge; limits; liquidity buffer','Daily price/basis/Greeks/VaR; stress when volatility spikes'),
('Credit','Due diligence; guaranteed performance; exposure caps','Before contracting and periodic reviews; recheck on adverse events'),
('Operational','Pre-trade validations; audit trail; reconciliation','Real-time alerts, daily exception review, incident escalation')],[1.1,2.75,2.6])

add_heading(d,'Methodological caveats',1)
d.add_paragraph('All figures are educational assumptions from the simulation, not market quotes, executable prices or assessed client exposures. The cost-of-carry expression assumes proportional storage costs; convenience yield is excluded. Black-76 is appropriate to an option on futures under its modeling assumptions. A daily or monthly VaR is not a maximum attainable loss. Real implementation requires contract specifications, basis analysis, execution/slippage, margin stress, legal review, counterparty assessment and approval against institutional policy.')
d.add_paragraph('Conclusion: Pair the appropriate hedge for the economic exposure with independent credit safeguards and automated execution controls. Evaluate all three risks together so that lowering price risk does not inadvertently create unacceptable liquidity, counterparty or operational risks.')

# keep titles attached to following paragraph
for para in d.paragraphs:
    if para.style.name.startswith('Heading'):para.paragraph_format.keep_with_next=True

d.core_properties.title='Coffee Futures - Risk Management Analysis';d.core_properties.subject='Citi MQA simulation - three risk scenarios';d.core_properties.author='Quantitative Analysis - Training Simulation'
d.save(out)
print('CREATED',out,'size',out.stat().st_size,'F=',F,'call=',call,'VaR=',var)
