from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from math import exp, log, sqrt
from scipy.stats import norm
import numpy as np
from pathlib import Path
OUT=Path('/mnt/data/Citi_Coffee_Structured_Note_Proposal.docx')
BLUE=RGBColor(22,57,100); DARK=RGBColor(29,43,61); GREEN=RGBColor(16,113,110); LIGHT=RGBColor(236,243,248); GREY=RGBColor(90,107,122)
S0=1.2;r=.02;d=.01;T=.5;sig=.25;K=1.25;N=100;c=.002;alpha=.10
F=S0*exp((r+d)*T)
d1=(log(F/K)+.5*sig*sig*T)/(sig*sqrt(T)); d2=d1-sig*sqrt(T)
call=exp(-r*T)*(F*norm.cdf(d1)-K*norm.cdf(d2))
principal_pv=N*exp(-r*T); coupon_pv=N*c*exp(-r*T); option_pv=N*alpha/K*call
headroom=N-principal_pv-coupon_pv-option_pv
rng=np.random.default_rng(42); n=200000
z=rng.standard_normal(n)
ST=S0*np.exp((r+d-.5*sig**2)*T + sig*np.sqrt(T)*z)
mc_pvs=np.exp(-r*T)*np.maximum(ST-K,0)
mc=mc_pvs.mean(); se=mc_pvs.std(ddof=1)/np.sqrt(n)

def set_cell_bg(cell,fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcpr.append(shd)
def set_cell_margin(cell,top=80,start=100,bottom=80,end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for n,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=OxmlElement(f'w:{n}');e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');m.append(e)
    tcPr.append(m)

doc=Document(); sec=doc.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.79);sec.right_margin=Inches(.79);sec.header_distance=Inches(.31);sec.footer_distance=Inches(.35)
styles=doc.styles
normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9);normal.font.color.rgb=DARK;normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.12
for name,size,color,spacebefore,spaceafter in [('Title',22,BLUE,0,9),('Heading 1',12.5,BLUE,12,5),('Heading 2',10,GREEN,8,3)]:
    st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=color;st.font.bold=True; st.paragraph_format.space_before=Pt(spacebefore);st.paragraph_format.space_after=Pt(spaceafter)
if 'Deck' not in styles:
 st=styles.add_style('Deck',WD_STYLE_TYPE.PARAGRAPH);st.base_style=normal;st.font.size=Pt(10);st.font.color.rgb=GREY;st.paragraph_format.space_after=Pt(10)
if 'SmallNote' not in styles:
 st=styles.add_style('SmallNote',WD_STYLE_TYPE.PARAGRAPH);st.base_style=normal;st.font.size=Pt(7.7);st.font.color.rgb=GREY;st.paragraph_format.space_after=Pt(4)
# header
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.LEFT
run=header.add_run('MARKETS QUANTITATIVE ANALYSIS   /   CITI TRAINING CASE');run.bold=True;run.font.size=Pt(7);run.font.color.rgb=BLUE
# footer
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('INDICATIVE • NOT AN INVESTMENT OFFER   |   ');run.font.size=Pt(7);run.font.color.rgb=GREY
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def para(s='',style=None):return doc.add_paragraph(s,style=style)
def heading(s,lev=1):doc.add_heading(s,level=lev)
def bullet(s):
 p=doc.add_paragraph(style='Normal');p.style='List Bullet';p.add_run(s);return p

def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):
  cell=t.rows[0].cells[i];set_cell_bg(cell,'163964');set_cell_margin(cell)
  run=cell.paragraphs[0].add_run(h);run.bold=True;run.font.size=Pt(8);run.font.color.rgb=RGBColor(255,255,255)
 for ri,row in enumerate(rows):
  cs=t.add_row().cells
  for i,val in enumerate(row):
   set_cell_margin(cs[i]);cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ri%2==0:set_cell_bg(cs[i],'F0F5F9')
   p=cs[i].paragraphs[0];p.paragraph_format.space_after=Pt(0)
   run=p.add_run(str(val));run.font.size=Pt(8)
   if i==0:run.bold=True
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths):row.cells[i].width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

# page 1
p=doc.add_paragraph(style='Title');p.add_run('Coffee-Linked Structured Note')
para('Structured Securities Proposal & Risk Mitigation Report  |  Scenario 2: Moderate-Risk Client','Deck')
heading('Executive recommendation')
para('Recommend an illustrative six-month, cash-settled coffee-linked note combining a small fixed coupon, principal repayment at maturity and limited participation in coffee price increases above a defined strike. The design responds to the client’s need for income and some commodity exposure without a direct position in futures.')
para('This is a preliminary training-case structure, not a tradable quote. Principal repayment depends on the issuer’s ability to pay; secondary-market prices can fall below par.')
heading('Proposed terms')
table(['Term','Illustrative specification'],[
 ('Investor profile','Moderate risk; seeks income and positive coffee participation'),
 ('Initial investment / maturity','$100 note denomination / six months'),
 ('Underlying / currency','Coffee reference price, USD per pound; cash settlement'),
 ('Reference level / strike','Initial spot $1.20/lb; upside strike $1.25/lb'),
 ('Fixed coupon','0.20% total over six months, paid at maturity'),
 ('Upside participation','10% of coffee appreciation above the $1.25/lb strike'),
 ('Principal protection','100% contractual principal at maturity, subject to issuer credit'),
 ('Observation / settlement','One fixing at maturity; detailed market conventions to be set')
],[1.95,4.9])
heading('Contractual payoff and investor outcomes')
para('For each $100 invested, maturity payment = $100 × [1 + 0.002 + 0.10 × max((C_T − 1.25)/1.25, 0)]. Here C_T is the defined coffee reference fixing (USD/lb) at the six-month maturity.')
table(['Coffee at maturity','Investor payout / $100','Interpretation'],[
 ('$0.95/lb','$100.20','Coffee falls: coupon and principal only'),
 ('$1.20/lb','$100.20','Unchanged versus initial spot: coupon only'),
 ('$1.25/lb','$100.20','At strike: no upside participation'),
 ('$1.40/lb','$101.40','12% rise above strike × 10% participation'),
 ('$1.60/lb','$103.00','28% rise above strike × 10% participation')
],[1.65,1.75,3.45])
para('Contractual payments exclude fees and taxes. No issuer-default recovery is assumed in this table.','SmallNote')

doc.add_page_break()
# page 2
heading('Quantitative valuation: models from the previous task')
para('Reuse Task 2 inputs as a controlled example: spot $1.20/lb, risk-free rate 2% annually, storage carry 1% annually, volatility 25% annually, and six-month maturity. These are training assumptions, not current market quotations.')
heading('1   Cost-of-carry futures benchmark',2)
para('F_0 = S_0 × exp[(r + d) × T] = 1.20 × exp[(0.02 + 0.01) × 0.5] = $1.21814/lb.')
para('The simple model assumes continuously compounded financing/storage rates and zero convenience yield. Before trading, calibrate delivery grade, location, inventories, convenience yield, forward curve and any spot-to-futures basis.')
heading('2   Option component and indicative funding budget',2)
para('For a European payoff on a futures-related coffee reference, use a Black-76-style formula: Call = exp(−rT)[F_0 N(d1) − K N(d2)], with d1 = [ln(F_0/K) + 0.5σ²T]/(σ√T), d2 = d1 − σ√T, and K = $1.25/lb.')
para(f'The theoretical call value is ${call:.5f}/lb. The note embeds 10% participation in (C_T − K)/K, equivalent to a call budget of $100 × 0.10/K × ${call:.5f} = ${option_pv:.4f} per $100 of principal.')
table(['Indicative cost at issue / $100','Present value'],[
 ('Risk-free funding of principal at maturity',f'${principal_pv:.4f}'),
 ('PV of $0.20 coupon at maturity',f'${coupon_pv:.4f}'),
 ('PV of embedded coffee call participation',f'${option_pv:.4f}'),
 ('Total theoretical cost',f'${principal_pv+coupon_pv+option_pv:.4f}'),
 ('Unallocated headroom before all costs',f'${headroom:.4f}')
],[4.35,2.5])
para('Headroom is only an indicative model budget, not an issuer margin: transaction fees, market-making hedges, issuer funding/credit, model reserves and slippage must be priced before issuance.','SmallNote')
heading('3   Monte Carlo validation and scenario analysis',2)
para('Simulate 200,000 independent six-month terminal coffee prices with geometric Brownian motion under a risk-neutral drift matching the simplified forward: C_T = S_0 × exp[(r + d − 0.5σ²)T + σ√T × Z], Z ~ N(0,1); random seed 42.')
para(f'Monte Carlo call estimate = ${mc:.5f}/lb (95% simulation confidence interval ${mc-1.96*se:.5f}–${mc+1.96*se:.5f}/lb), consistent with the analytical ${call:.5f}/lb value. This is a model consistency check, not an empirical forecast.')
para('For supply-demand, climate and geopolitical stresses, perturb the forward curve and volatility separately. Strong harvests may lower expected prices; crop damage, export restrictions and logistics shocks may raise prices and volatility. No event probabilities are implied by these directional scenarios.')

# page3

doc.add_page_break()
heading('Hedging implementation and risk mitigation')
para('Issuer hedge: buy a matching coffee call in the market where available, or dynamically delta-hedge the short embedded call exposure with the appropriate futures contract. Fund the protected payment with high-quality discounted assets matched to maturity, subject to funding and counterparty limits.')
table(['Risk','Mitigation and control'],[
 ('Coffee price / delta','Daily option Greeks; set delta/vega limits, rebalance hedge and stress large gaps.'),
 ('Basis, delivery and roll','Match ICE grade, reference fixing, expiry and units; monitor spot–futures basis and roll.'),
 ('Volatility / model risk','Compare implied and historical volatility; reprice with Black-76 and Monte Carlo; independent validation.'),
 ('Liquidity / execution','Limit position size, maintain liquid hedging instruments and assess bid-ask spreads.'),
 ('Issuer / counterparty credit','Disclose that protection is contractual, not insured; set exposure limits and review credit quality.'),
 ('Financing / interest rates','Match funded liabilities and maturities; monitor discounting and funding spread changes.'),
 ('Operations / legal','Daily reconciliation; confirm documentation, fixings, settlement and investor disclosures.')
],[1.67,5.18])
heading('Key sensitivity: volatility can exhaust pricing headroom')
def opt_price(vol):
 a=(log(F/K)+.5*vol**2*T)/(vol*sqrt(T));b=a-vol*sqrt(T)
 return exp(-r*T)*(F*norm.cdf(a)-K*norm.cdf(b))
rows=[]
for v in [.15,.25,.35]:
 cp=opt_price(v)
 cost=N*alpha/K*cp
 head=N-principal_pv-coupon_pv-cost
 rows.append((f'{v:.0%}',f'${cp:.5f}/lb',f'${cost:.4f}',f'${head:+.4f}'))
table(['Implied volatility','Call fair value','Note option cost / $100','Theoretical headroom'],rows,[1.45,1.55,1.8,2.05])
para('At 35% volatility, the structure becomes theoretically underfunded at the same offered terms. Reprice participation/coupon, reduce exposure, obtain executable hedges or decline issuance; never assume a historical-volatility estimate is an executable option quote.')
heading('Decision and next steps')
para('Proceed only to an indicative client discussion. Before issuance: (i) obtain the relevant coffee futures and option quotes, (ii) confirm legal reference/settlement terms and suitability, (iii) validate funding and hedge costs, (iv) run downside, gap, volatility and default stresses, and (v) secure independent model and risk approval.')
para('Provenance: client scenarios and product techniques are supplied by Citi/Forage Task 3; price inputs and core models are reused from the earlier Task 2 learning exercise. The coupon, strike use, participation rate, structure, risk limits and Monte Carlo valuation here are analyst-designed illustrative assumptions, not statements about Citi products.','SmallNote')

# set page-level keep with next for headings
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):
  p.paragraph_format.keep_with_next=True
# Save
doc.save(OUT)
print('Saved:',OUT)
print('CALCS',{'future':F,'call':call,'principal_pv':principal_pv,'coupon_pv':coupon_pv,'option_pv':option_pv,'headroom':headroom,'mc':mc,'mc_ci':(mc-1.96*se,mc+1.96*se)})
